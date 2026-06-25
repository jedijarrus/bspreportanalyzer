"""Eigenständige Explorations-Instanz — UNABHÄNGIG vom App-Code.

Importiert alle echten xlsx aus data/uploads/ in eine separate SQLite
(data/explore.db, gitignored) und erlaubt PII-sichere Aggregat-Auswertung.
Dient nur dem Verstehen der Daten — nicht Teil der Anwendung.

    python scripts/explore.py import      # xlsx -> data/explore.db
    python scripts/explore.py profile     # aggregiertes Profil (keine PII)
    python scripts/explore.py sql "SELECT ..."   # freie Query (Vorsicht: PII!)

Bewusst ohne Abhängigkeit zu app/ — eigene Mini-Logik, damit Exploration
und Produktcode getrennt bleiben.
"""
from __future__ import annotations

import argparse
import datetime as dt
import re
import sqlite3
import sys
import unicodedata
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
UPLOADS = ROOT / "data" / "uploads"
DB = ROOT / "data" / "explore.db"
SHEET_FALLBACK = "Kundennummer"

# Felder, die als PII gelten — werden in profile() NIE im Klartext gezeigt.
PII_FIELDS = {
    "gp_firmenname", "gp_namenszusatz", "gp_nachname", "gp_vorname", "gp_strasse",
    "gp_wohnort", "gp_plz", "gp_adresszusatz", "re_firmenname", "re_strasse",
    "re_wohnort", "re_plz", "re_nachname", "re_vorname", "rufnummer", "iban", "bic",
    "kreditinstitut", "karten_profilnummer", "eid", "kundennummer", "kundenkonto",
    "kostenstelle", "kostenstellennutzer", "auftragsnummer", "rufnummer_combicard",
    "combicard_profilnummer", "rahmenvertrag",
}


def slug(header: str) -> str:
    s = header.strip().lower()
    s = s.replace("ä", "ae").replace("ö", "oe").replace("ü", "ue").replace("ß", "ss")
    s = unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode("ascii")
    return re.sub(r"[^a-z0-9]+", "_", s).strip("_")


def _cell(v):
    if isinstance(v, dt.datetime):
        return v.isoformat()
    return v


def cmd_import() -> None:
    files = sorted(UPLOADS.glob("*.xlsx"))
    if not files:
        print(f"keine xlsx in {UPLOADS}")
        return
    if DB.exists():
        DB.unlink()
    con = sqlite3.connect(DB)

    headers = None
    for fi, f in enumerate(files):
        ws = openpyxl.load_workbook(f, read_only=True, data_only=True).active
        rows = list(ws.iter_rows(values_only=True))
        cols = [slug(str(h)) for h in rows[0]]
        if headers is None:
            headers = cols
            coldef = ", ".join(f'"{c}" TEXT' for c in headers)
            con.execute(
                f'CREATE TABLE c (_src TEXT, _report_date TEXT, {coldef})'
            )
        # Report-Datum aus Dateiname: 20260625-084350_...
        m = re.match(r"(\d{8})-(\d{6})", f.name)
        rdate = (
            dt.datetime.strptime(m.group(1) + m.group(2), "%Y%m%d%H%M%S").isoformat()
            if m else None
        )
        body = [r for r in rows[1:] if r[1] is not None]
        placeholders = ", ".join("?" * (len(headers) + 2))
        con.executemany(
            f'INSERT INTO c VALUES ({placeholders})',
            [(f.name, rdate, *[_cell(v) for v in r[: len(headers)]]) for r in body],
        )
        print(f"  importiert: {f.name}  ({len(body)} Zeilen)")
    con.commit()
    n = con.execute("SELECT COUNT(*) FROM c").fetchone()[0]
    print(f"-> {DB.name}: {n} Zeilen gesamt, {len(headers)} Spalten")
    con.close()


def _con() -> sqlite3.Connection:
    if not DB.exists():
        print("erst importieren: python scripts/explore.py import")
        sys.exit(1)
    return sqlite3.connect(DB)


def cmd_profile() -> None:
    con = _con()
    cols = [r[1] for r in con.execute("PRAGMA table_info(c)")][2:]
    N = con.execute("SELECT COUNT(*) FROM c").fetchone()[0]
    print("=" * 60)
    print(f"AGGREGAT-PROFIL  N={N} Linien  (PII-sicher)")
    print("=" * 60)

    print("\n### Reports in der DB")
    for src, rd, cnt in con.execute(
        "SELECT _src, _report_date, COUNT(*) FROM c GROUP BY _src ORDER BY _report_date"
    ):
        print(f"  {rd[:10] if rd else '?'}  {cnt:4} Linien  ({src[:30]})")

    print("\n### Spalten: leer / konstant / variabel")
    empty, const, var = [], [], []
    for col in cols:
        ne = con.execute(
            f'SELECT COUNT(*), COUNT(DISTINCT "{col}") FROM c '
            f'WHERE "{col}" IS NOT NULL AND "{col}" != ""'
        ).fetchone()
        filled, distinct = ne
        if filled == 0:
            empty.append(col)
        elif distinct == 1:
            const.append(col)
        else:
            var.append((col, distinct, N - filled))
    print(f"  komplett leer : {len(empty)}")
    print(f"  konstant      : {len(const)}  (1 Wert über alle Zeilen)")
    print(f"  variabel      : {len(var)}")
    print("\n  -- variable Spalten [distinct | leer] --")
    for col, d, e in sorted(var, key=lambda x: -x[1]):
        pii = "  [PII]" if col in PII_FIELDS else ""
        print(f"     {col:34} d={d:<4} leer={e}{pii}")

    print("\n### Kategorische Verteilungen (nicht-PII)")
    for col in ["kartentyp", "tarif", "vertragsstatus", "daten_optionen",
                "voice_optionen", "roaming_optionen", "vvl_berechtigung",
                "rechnungszahlart", "evn"]:
        if col not in cols:
            continue
        rows = con.execute(
            f'SELECT COALESCE(NULLIF("{col}",\'\'),\'(leer)\'), COUNT(*) '
            f'FROM c GROUP BY 1 ORDER BY 2 DESC LIMIT 8'
        ).fetchall()
        print(f"  {col}:")
        for v, c in rows:
            print(f"     {c:4}  {v}")

    print("\n### Sperren (gesperrte Linien)")
    locked = con.execute(
        "SELECT COUNT(*) FROM c WHERE sperren IS NOT NULL AND sperren != ''"
    ).fetchone()[0]
    print(f"  gesperrt: {locked} / {N}")

    print("\n### Bindefristende relativ zu heute")
    today = dt.date(2026, 6, 25).isoformat()
    buckets = {
        "abgelaufen": ('bindefristende < ?', (today,)),
        "0-3 Monate": ('bindefristende >= ? AND bindefristende < ?', (today, dt.date(2026, 9, 25).isoformat())),
        "3-12 Monate": ('bindefristende >= ? AND bindefristende < ?', (dt.date(2026, 9, 25).isoformat(), dt.date(2027, 6, 25).isoformat())),
        ">12 Monate": ('bindefristende >= ?', (dt.date(2027, 6, 25).isoformat(),)),
    }
    for label, (cond, params) in buckets.items():
        c = con.execute(f"SELECT COUNT(*) FROM c WHERE {cond}", params).fetchone()[0]
        print(f"  {label:14}: {c}")

    print("\n### Vertragsbeginn nach Jahr")
    for yr, c in con.execute(
        "SELECT substr(vertragsbeginn,1,4) y, COUNT(*) FROM c "
        "WHERE vertragsbeginn IS NOT NULL GROUP BY y ORDER BY y"
    ):
        print(f"  {yr}: {c}")

    print("\n### MultiSIM-Nutzung")
    ms = con.execute(
        'SELECT COUNT(*) FROM c WHERE "multisim_karten_profilnummer_1" IS NOT NULL '
        'AND "multisim_karten_profilnummer_1" NOT IN (\'\',\'\\\\\')'
    ).fetchone()[0]
    print(f"  Linien mit MultiSIM 1: {ms} / {N}")
    con.close()


def cmd_sql(query: str) -> None:
    con = _con()
    for row in con.execute(query):
        print(row)
    con.close()


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__)
    sub = ap.add_subparsers(dest="cmd", required=True)
    sub.add_parser("import")
    sub.add_parser("profile")
    p = sub.add_parser("sql"); p.add_argument("query")
    args = ap.parse_args()
    {"import": cmd_import, "profile": cmd_profile,
     "sql": lambda: cmd_sql(args.query)}[args.cmd]()


if __name__ == "__main__":
    main()
